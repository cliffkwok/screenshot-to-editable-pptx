"""ADK agent workflow — Part 1: setup + frame + helpers"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

WHITE  = RGBColor(0xFF,0xFF,0xFF)
BLACK  = RGBColor(0x00,0x00,0x00)
GRAY   = RGBColor(0x5F,0x63,0x68)
CARD_BR = RGBColor(0xE0,0xE0,0xE0)
TEXT   = RGBColor(0x09,0x09,0x09)

# Color themes (hex from vision analysis)
BLUE_H = RGBColor(0x45,0x83,0xEC)
BLUE_D = RGBColor(0x1A,0x56,0xDB)
YELLOW = RGBColor(0xF5,0xC1,0x44)
YELLOW_H=RGBColor(0xFB,0xBC,0x04)
YELLOW_D=RGBColor(0xE8,0x71,0x0A)
RED    = RGBColor(0xDD,0x3A,0x33)
RED_H  = RGBColor(0xEA,0x99,0x99)
RED_D  = RGBColor(0xC5,0x22,0x1F)

def S(typ,l,t,w,h,fill=None,line=None,lw=None,rad=None):
    s=slide.shapes.add_shape(typ,l,t,w,h)
    if fill: s.fill.solid(); s.fill.fore_color.rgb=fill
    else: s.fill.background()
    if line: s.line.color.rgb=line; s.line.width=lw or Pt(1)
    else: s.line.fill.background()
    if rad is not None: s.adjustments[0]=rad
    s.shadow.inherit = False
    return s

def T(l,t,w,h,txt,sz,bold=False,c=TEXT,al=PP_ALIGN.CENTER):
    tb=slide.shapes.add_textbox(l,t,w,h)
    tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=txt
    p.font.size=sz; p.font.bold=bold; p.font.color.rgb=c
    p.font.name="Poppins"; p.alignment=al
    return tb

# Main slide frame
FX, FY = Inches(0.3), Inches(0.3)
FW, FH = Inches(12.73), Inches(6.9)
S(MSO_SHAPE.ROUNDED_RECTANGLE, FX, FY, FW, FH, line=BLACK, lw=Pt(2), rad=0.02)

def P(xp,yp,wp,hp):
    """Convert % to Inches relative to frame."""
    return FX+FW*xp/100, FY+FH*yp/100, FW*wp/100, FH*hp/100

# Title
x,y,w,h = P(30,4,40,6)
T(x,y,w,h,"ADK agent - workflow", Pt(24), bold=True)


def group_shapes(name, shapes):
    """Group list of shapes together using XML."""
    if len(shapes) < 2:
        return
    from lxml import etree
    from pptx.oxml.ns import qn
    sp_tree = shapes[0]._element.getparent()
    grp = etree.SubElement(sp_tree, qn('p:grpSp'))
    nv = etree.SubElement(grp, qn('p:nvGrpSpPr'))
    cnv = etree.SubElement(nv, qn('p:cNvPr'), {'id': '0', 'name': name})
    cnv_grp = etree.SubElement(nv, qn('p:cNvGrpSpPr'))
    nv_pr = etree.SubElement(nv, qn('p:nvPr'))
    grp_pr = etree.SubElement(grp, qn('p:grpSpPr'))
    xfrm = etree.SubElement(grp_pr, qn('a:xfrm'))
    for s in shapes:
        sp_tree.remove(s._element)
        grp.append(s._element)
    print(f"Group '{name}': {len(shapes)} shapes")

def card_circle(cx,cy,r,fill_color,label):
    S(MSO_SHAPE.OVAL, cx-r, cy-r, r*2, r*2, fill=fill_color)
    T(cx-r*2, cy-r*0.35, r*4, r*0.7, label, Pt(6), bold=True, c=WHITE)
print("Part 1 done — setup")"""Part 2 — Top row card: Sequential (blue)"""

def make_top_card(x,y,w,h,header_text,header_bg,footer_text):
    """Create a top-row card with header bar + content area."""
    # Card background
    S(MSO_SHAPE.RECTANGLE, x, y, w, h, fill=WHITE, line=GRAY, lw=Pt(1))
    # Header bar
    S(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.35), fill=header_bg)
    T(x, y, w, Inches(0.35), header_text, Pt(12), bold=True, c=WHITE)
    # Footer text
    fy = y + h - Inches(0.35)
    T(x+Inches(0.1), fy, w-Inches(0.2), Inches(0.35), footer_text, Pt(8), c=GRAY)
    return y+Inches(0.35), fy  # content area top, bottom


# Sequential card position: x=4-31%, y=20-47%
x,y,w,h = P(4,20,27,27)
ct, cb = make_top_card(x,y,w,h,"Sequential", BLUE_H, "Chains multiple specialized LLMs in a pipeline.")

# Content area
cy_ctr = (ct + cb) / 2
# Circles: Input (blue) → LLM (dark) → LLM (dark) → Output (blue)
cx1 = x + w*0.12; cx2 = x + w*0.38; cx3 = x + w*0.62; cx4 = x + w*0.88
r = Inches(0.28)
card_circle(cx1, cy_ctr, r, BLUE_H, "Input")
card_circle(cx2, cy_ctr, r, BLUE_D, "LLM")
card_circle(cx3, cy_ctr, r, BLUE_D, "LLM")
card_circle(cx4, cy_ctr, r, BLUE_H, "Output")

# Connector lines (thin rectangles)
cx = cx1+r; ex = cx2-r
S(MSO_SHAPE.RECTANGLE, cx, cy_ctr-Pt(0.5), ex-cx, Pt(1), fill=GRAY)
cx = cx2+r; ex = cx3-r
S(MSO_SHAPE.RECTANGLE, cx, cy_ctr-Pt(0.5), ex-cx, Pt(1), fill=GRAY)
cx = cx3+r; ex = cx4-r
S(MSO_SHAPE.RECTANGLE, cx, cy_ctr-Pt(0.5), ex-cx, Pt(1), fill=GRAY)

# Labels above/below
T(cx1+r-Inches(0.5), ct+Inches(0.1), cx3-cx1-Inches(0.1), Inches(0.25), "Intermediate result...", Pt(7), c=GRAY)
T(cx2+r-Inches(0.5), cb-Inches(0.5), cx4-cx2-Inches(0.1), Inches(0.25), "Intermediate Result...", Pt(7), c=GRAY)

"""ADK agent workflow — remaining cards"""

# Parallel card (yellow) x=33-60%, y=20-47%
x,y,w,h = P(33,20,27,27)
S(MSO_SHAPE.RECTANGLE, x, y, w, h, fill=WHITE, line=GRAY, lw=Pt(1))
S(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.35), fill=YELLOW)
T(x, y, w, Inches(0.35), "Parallel", Pt(12), bold=True, c=WHITE)
fy = y+h-Inches(0.35)
T(x+Inches(0.1), fy, w-Inches(0.2), Inches(0.35), "Multiple LLMs concurrently and aggregate outputs", Pt(8), c=GRAY)
cy = y+Inches(0.35)+(fy-y-Inches(0.35))/2
# Input → 3 LLM → Aggregator → Output
cx1=x+w*0.08; cx2=x+w*0.42; cx3=x+w*0.72; cx4=x+w*0.92; r=Inches(0.28)
card_circle(cx1, cy, r, YELLOW_H, "Input")
S(MSO_SHAPE.OVAL, cx2-r, y+Inches(0.35)+(fy-y-Inches(0.35))*0.15, r*2, r*2, fill=YELLOW_D)
T(cx2-3*r, y+Inches(0.35)+(fy-y-Inches(0.35))*0.15-r, r*6, r*1.5, "LLM", Pt(7), bold=True, c=WHITE)
S(MSO_SHAPE.OVAL, cx2-r, cy-r, r*2, r*2, fill=YELLOW_D)
T(cx2-3*r, cy-r*1.5, r*6, r*1.5, "LLM", Pt(7), bold=True, c=WHITE)
S(MSO_SHAPE.OVAL, cx2-r, y+Inches(0.35)+(fy-y-Inches(0.35))*0.85-r, r*2, r*2, fill=YELLOW_D)
T(cx2-3*r, y+Inches(0.35)+(fy-y-Inches(0.35))*0.85-r*2, r*6, r*1.5, "LLM", Pt(7), bold=True, c=WHITE)
card_circle(cx3, cy, r, YELLOW_H, "Aggregator")
card_circle(cx4, cy, r, YELLOW_H, "Output")
# Lines: Input→3 LLM and 3 LLM→Aggregator→Output
for cy_llm in [y+Inches(0.35)+(fy-y-Inches(0.35))*0.15, cy, y+Inches(0.35)+(fy-y-Inches(0.35))*0.85]:
    S(MSO_SHAPE.RECTANGLE, cx1+r, cy_llm-Pt(0.5), cx2-r-cx1-r, Pt(1), fill=GRAY)
    S(MSO_SHAPE.RECTANGLE, cx2+r, cy_llm-Pt(0.5), cx3-r-cx2-r, Pt(1), fill=GRAY)
S(MSO_SHAPE.RECTANGLE, cx3+r, cy-Pt(0.5), cx4-r-cx3-r, Pt(1), fill=GRAY)



# Hierarchical card (red) x=62-99%, y=20-47%
x,y,w,h = P(62,20,37,27)
S(MSO_SHAPE.RECTANGLE, x, y, w, h, fill=WHITE, line=GRAY, lw=Pt(1))
S(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.35), fill=RED)
T(x, y, w, Inches(0.35), "Hierarchical", Pt(12), bold=True, c=WHITE)
fy = y+h-Inches(0.35)
T(x+Inches(0.1), fy, w-Inches(0.2), Inches(0.35), "Router model to direct inputs to specialized LLMs based on content or intent", Pt(8), c=GRAY)
cy = y+Inches(0.35)+(fy-y-Inches(0.35))/2
cx1=x+w*0.07; cx2=x+w*0.40; cx3=x+w*0.88; r=Inches(0.2)
# Input circle
card_circle(cx1, cy, r, RED, "Input")
# Router circle (pink)
card_circle(cx2, y+Inches(0.40)+fy-y-Inches(0.40)*0.12, r*1.2, RED_H, "LLM")
T(cx2-3*r, y+Inches(0.40)+fy-y-Inches(0.40)*0.12-r*2, r*6, r*2, "processing\nRouter", Pt(7), bold=True, c=TEXT)
# Analysis text
T(cx2-3*r, y+Inches(0.40)+fy-y-Inches(0.40)*0.12+r*1.5, r*4, r*1.5, "Analysis and classification", Pt(7), c=GRAY)
# 3 LLM circles (dark red)
offsets = [-Inches(0.6), 0, Inches(0.6)]
for off in offsets:
    S(MSO_SHAPE.OVAL, cx2-off-r, y+Inches(0.40)+fy-y-Inches(0.40)*0.68, r*2, r*2, fill=RED_D)
    T(cx2-off-2*r, y+Inches(0.40)+fy-y-Inches(0.40)*0.68-r*1.5, r*4, r*1.5, "LLM", Pt(7), bold=True, c=WHITE)
# Output circle
card_circle(cx3, cy, r*1.2, RED, "Output")
# Lines from Router to LLMs
for off in offsets:
    S(MSO_SHAPE.RECTANGLE, cx2-off, cy-Pt(0.5), Pt(1), y+Inches(0.40)+fy-y-Inches(0.40)*0.68-cy, fill=GRAY)
# Lines from LLMs to Output
for off in offsets:
    lx = cx2-off+r; ex = cx3-r
    S(MSO_SHAPE.RECTANGLE, lx, y+Inches(0.40)+fy-y-Inches(0.40)*0.68+r*2+Pt(1), ex-lx, Pt(1), fill=GRAY)
# Input→Router line
S(MSO_SHAPE.RECTANGLE, cx1+r, cy-Pt(0.5), cx2-r-cx1-r, Pt(1), fill=GRAY)
S(MSO_SHAPE.RECTANGLE, cx2+r, y+h*0.72, cx3-r-cx2-r, Pt(1), fill=GRAY)

def bottom_card(x,y,w,h,title,accent):
    S(MSO_SHAPE.RECTANGLE,x,y,w,h,fill=WHITE,line=CARD_BR,lw=Pt(1))
    T(x,y,w,Inches(0.3),title,Pt(11),bold=True,c=BLACK)
    S(MSO_SHAPE.RECTANGLE,x,y+Inches(0.3),w,Pt(1.5),fill=accent)
    return y+Inches(0.35),y+h
def pill(x,y,accent):
    pw,ph=Inches(0.5),Inches(0.32)
    S(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,pw,ph,fill=accent,rad=0.5)
    S(MSO_SHAPE.OVAL,x+Inches(.13),y+Inches(.08),Inches(.07),Inches(.07),fill=WHITE)
    S(MSO_SHAPE.OVAL,x+Inches(.30),y+Inches(.08),Inches(.07),Inches(.07),fill=WHITE)
    return pw,ph
# 1. Sequential agent (blue)
x,y,w,h = P(4,50,26,47)
ct,cb = bottom_card(x,y,w,h,"Sequential agent",BLUE_H)
cy=(ct+cb)/2; ph_val=Inches(0.32)
for i in range(3):
    px=x+w*0.10+i*w*0.30; py=cy-ph_val/2
    pw,_=pill(px,py,BLUE_H)
    if i<2: S(MSO_SHAPE.RECTANGLE,px+pw,cy-Pt(0.5),x+w*0.10+(i+1)*w*0.30-px-pw,Pt(1.5),fill=GRAY)

def arrow_h(x1,x2,cy):
    S(MSO_SHAPE.RECTANGLE,x1,cy-Pt(0.5),x2-x1,Pt(1.5),fill=GRAY)

# 2. Parallel agent (yellow)
x,y,w,h = P(32,50,25,47)
ct,cb = bottom_card(x,y,w,h,"Parallel agent",YELLOW)
cy=(ct+cb)/2; ph=Inches(0.32)
for i in range(3):
    py=ct+(cb-ct)*0.15+i*(cb-ct)*0.28
    pill(x+w*0.35,py,YELLOW)
    S(MSO_SHAPE.RECTANGLE,x+w*0.05,py+ph/2-Pt(0.5),x+w*0.35-x-w*0.05,Pt(1.5),fill=GRAY)

# 3. Loop agent (red)
x,y,w,h = P(59,50,26,47)
ct,cb = bottom_card(x,y,w,h,"Loop agent",RED)
cy=(ct+cb)/2
pos=[(x+w*0.20,cy-Inches(.3)),(x+w*0.55,cy-Inches(.3)),
     (x+w*0.20,cy+Inches(.3)),(x+w*0.55,cy+Inches(.3))]
for px,py in pos: pill(px,py,RED)
# arrows connecting loop
arrow_h(x+w*0.02,pos[0][0],pos[0][1]+Inches(.16))
arrow_h(pos[0][0]+Inches(.5),pos[1][0],pos[0][1]+Inches(.16))
S(MSO_SHAPE.RECTANGLE,pos[1][0]+Inches(.1),pos[1][1]+Inches(.32),Pt(1.5),pos[3][1]-pos[1][1]-Inches(.16),fill=GRAY)
arrow_h(pos[3][0]+Inches(.5),pos[2][0]+Inches(.5),pos[3][1]+Inches(.16))
S(MSO_SHAPE.RECTANGLE,pos[2][0]+Inches(.1),pos[0][1]+Inches(.32),Pt(1.5),pos[2][1]-pos[0][1]-Inches(.16),fill=GRAY)

# 4. LLM agent (red)
x,y,w,h = P(87,50,12,47)
ct,cb = bottom_card(x,y,w,h,"LLM agent",RED)
py_bot=cb-Inches(.55)
ps=[(x+w*0.08,py_bot),(x+w*0.35,py_bot),(x+w*0.62,py_bot)]
for px,py in ps: pill(px,py,RED)
apex=x+w*0.42; ay=ct+Inches(.25)
for px,py in ps:
    S(MSO_SHAPE.RECTANGLE,apex,ay,Pt(1),py-ay,fill=GRAY)

