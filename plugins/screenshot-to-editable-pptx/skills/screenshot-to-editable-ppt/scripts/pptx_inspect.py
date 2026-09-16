#!/usr/bin/env python3
"""Dump every shape/text property from a PPTX slide to JSON.

Usage:
  python3 pptx_inspect.py --pptx slide.pptx --out inspect.json
  python3 pptx_inspect.py --pptx slide.pptx --slide 0 --out inspect.json
"""
from __future__ import annotations
import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


def hex_rgb(rgb) -> str | None:
    if rgb is None:
        return None
    try:
        return f"#{int(rgb[0]):02X}{int(rgb[1]):02X}{int(rgb[2]):02X}"
    except Exception:
        return None


def safe_fill(shape):
    try:
        fill = shape.fill
        if fill.type is None:
            return None
        return hex_rgb(fill.fore_color.rgb)
    except Exception:
        return None


def safe_line(shape):
    try:
        line = shape.line
        color = hex_rgb(line.color.rgb)
        width_pt = None
        if line.width is not None:
            width_pt = round(line.width.pt, 2)
        return {"color": color, "width_pt": width_pt}
    except Exception:
        return {"color": None, "width_pt": None}


def has_outer_shadow(shape) -> bool:
    el = shape._element
    for node in el.iter():
        if node.tag == qn("a:outerShdw"):
            return True
    return False


def auto_shape_name(shape) -> str | None:
    try:
        ast = shape.auto_shape_type
        return ast.name if ast is not None else None
    except Exception:
        return None


def radius_of(shape):
    try:
        return round(float(shape.adjustments[0]), 3)
    except Exception:
        return None


def _font_color(font):
    try:
        return hex_rgb(font.color.rgb)
    except Exception:
        return None


def _font_fields(font):
    size_pt = None
    try:
        if font.size is not None:
            size_pt = round(font.size.pt, 2)
    except Exception:
        pass
    bold = None
    try:
        bold = font.bold
    except Exception:
        pass
    return {
        "font_name": font.name,
        "font_size_pt": size_pt,
        "bold": bold,
        "color": _font_color(font),
    }


def _align_name(alignment):
    if alignment is None:
        return None
    name = str(alignment)
    if "CENTER" in name:
        return "center"
    if "RIGHT" in name:
        return "right"
    if "LEFT" in name:
        return "left"
    if "JUSTIFY" in name:
        return "justify"
    return name.lower()


def font_info(shape):
    if not getattr(shape, "has_text_frame", False):
        return None
    tf = shape.text_frame
    info = {"text": tf.text or "", "runs": []}
    for p in tf.paragraphs:
        pf = _font_fields(p.font)
        align = _align_name(p.alignment)
        if p.runs:
            for r in p.runs:
                rf = _font_fields(r.font)
                info["runs"].append({
                    "text": r.text,
                    "font_name": rf["font_name"] or pf["font_name"],
                    "font_size_pt": rf["font_size_pt"] if rf["font_size_pt"] is not None else pf["font_size_pt"],
                    "bold": rf["bold"] if rf["bold"] is not None else bool(pf["bold"]),
                    "color": rf["color"] or pf["color"],
                    "align": align,
                })
        elif (p.text or "").strip():
            info["runs"].append({
                "text": p.text,
                "font_name": pf["font_name"],
                "font_size_pt": pf["font_size_pt"],
                "bold": bool(pf["bold"]),
                "color": pf["color"],
                "align": align,
            })
    return info


def kind_of(shape) -> str:
    t = shape.shape_type
    if t == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if t == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if t == MSO_SHAPE_TYPE.TEXT_BOX:
        return "textbox"
    if t == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "shape"
    return str(t)


def _xfrm_of(el):
    for child in el:
        if child.tag in (qn("p:spPr"), qn("p:grpSpPr")):
            xf = child.find(qn("a:xfrm"))
            if xf is not None:
                return xf
    return el.find(".//" + qn("a:xfrm"))


def _xy(node, x="x", y="y"):
    if node is None:
        return 0, 0
    return int(float(node.get(x, 0))), int(float(node.get(y, 0)))


def read_group_xfrm(shape):
    """Raw OOXML group transform. PowerPoint uses this, not python-pptx .left."""
    xf = _xfrm_of(shape._element)
    off = xf.find(qn("a:off")) if xf is not None else None
    ext = xf.find(qn("a:ext")) if xf is not None else None
    ch_off = xf.find(qn("a:chOff")) if xf is not None else None
    ch_ext = xf.find(qn("a:chExt")) if xf is not None else None
    ox, oy = _xy(off)
    cx, cy = _xy(ext, "cx", "cy")
    cox, coy = _xy(ch_off)
    ccx, ccy = _xy(ch_ext, "cx", "cy")
    return {
        "off": (ox, oy),
        "ext": (cx, cy),
        "chOff": (cox, coy),
        "chExt": (ccx or 1, ccy or 1),
    }


def group_coord_mode(gx, child_offs):
    cox, coy = gx["chOff"]
    if cox == 0 and coy == 0:
        return "local"
    if not child_offs:
        return "absolute"
    ccx, ccy = gx["chExt"]
    localish = all(0 <= x <= ccx and 0 <= y <= ccy for x, y in child_offs)
    if localish:
        return "DOUBLE_OFFSET_BUG"
    return "absolute"


def ppt_world(local_x, local_y, gx):
    """PowerPoint world EMUs: off + (child.off - chOff) * (ext / chExt)."""
    sx = gx["ext"][0] / gx["chExt"][0] if gx["chExt"][0] else 1
    sy = gx["ext"][1] / gx["chExt"][1] if gx["chExt"][1] else 1
    wx = gx["off"][0] + (int(local_x) - gx["chOff"][0]) * sx
    wy = gx["off"][1] + (int(local_y) - gx["chOff"][1]) * sy
    return int(wx), int(wy)


def record_shape(shape, slide_w, slide_h, group_path="", ox=0, oy=0):
    left = int(shape.left) + int(ox)
    top = int(shape.top) + int(oy)
    width = int(shape.width)
    height = int(shape.height)
    rec = {
        "name": shape.name,
        "kind": kind_of(shape),
        "shape_type": auto_shape_name(shape),
        "group_path": group_path,
        "left_in": round(left / Inches(1), 4),
        "top_in": round(top / Inches(1), 4),
        "width_in": round(width / Inches(1), 4),
        "height_in": round(height / Inches(1), 4),
        "x_pct": round(left / slide_w * 100, 2),
        "y_pct": round(top / slide_h * 100, 2),
        "w_pct": round(width / slide_w * 100, 2),
        "h_pct": round(height / slide_h * 100, 2),
        "fill": safe_fill(shape),
        "line": safe_line(shape),
        "radius": radius_of(shape),
        "has_shadow": has_outer_shadow(shape),
        "font": font_info(shape),
    }
    if rec["kind"] == "shape" and rec["shape_type"] == "OVAL":
        rec["is_circle"] = abs(width - height) / max(width, height) < 0.05
    return rec


def walk(shapes, slide_w, slide_h, group_path="", mapper=None):
    """mapper(local_left, local_top) -> slide-absolute EMUs using PowerPoint's xfrm formula."""
    if mapper is None:
        mapper = lambda l, t: (int(l), int(t))
    out = []
    groups = 0
    pictures = 0
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            groups += 1
            path = f"{group_path}/{s.name}" if group_path else s.name
            gx = read_group_xfrm(s)
            child_offs = [(int(c.left), int(c.top)) for c in s.shapes]
            mode = group_coord_mode(gx, child_offs)
            sx = gx["ext"][0] / gx["chExt"][0] if gx["chExt"][0] else 1
            sy = gx["ext"][1] / gx["chExt"][1] if gx["chExt"][1] else 1

            def child_mapper(l, t, _gx=gx, _map=mapper, _sx=sx, _sy=sy):
                lx = _gx["off"][0] + (int(l) - _gx["chOff"][0]) * _sx
                ly = _gx["off"][1] + (int(t) - _gx["chOff"][1]) * _sy
                return _map(lx, ly)

            child = walk(s.shapes, slide_w, slide_h, path, child_mapper)
            out.extend(child["shapes"])
            groups += child["groups"]
            pictures += child["pictures"]
            gwx, gwy = mapper(gx["off"][0], gx["off"][1])
            out.append({
                "name": s.name,
                "kind": "group",
                "group_path": group_path,
                "child_count": len(list(s.shapes)),
                "coord_mode": mode,
                "chOff_emu": list(gx["chOff"]),
                "x_pct": round(gwx / slide_w * 100, 2),
                "y_pct": round(gwy / slide_h * 100, 2),
                "w_pct": round(gx["ext"][0] / slide_w * 100, 2),
                "h_pct": round(gx["ext"][1] / slide_h * 100, 2),
            })
        else:
            wx, wy = mapper(int(s.left), int(s.top))
            rec = record_shape(s, slide_w, slide_h, group_path,
                               ox=wx - int(s.left), oy=wy - int(s.top))
            if rec["kind"] == "picture":
                pictures += 1
            out.append(rec)
    return {"shapes": out, "groups": groups, "pictures": pictures}


def inspect_pptx(path: Path, slide_index=0):
    prs = Presentation(str(path))
    slide = prs.slides[slide_index]
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    bg = None
    try:
        bg = hex_rgb(slide.background.fill.fore_color.rgb)
    except Exception:
        pass
    walked = walk(slide.shapes, sw, sh)
    return {
        "file": str(path),
        "slide_index": slide_index,
        "slide": {
            "width_in": round(sw / Inches(1), 4),
            "height_in": round(sh / Inches(1), 4),
            "background": bg,
            "top_level_shapes": len(slide.shapes),
        },
        "groups": walked["groups"],
        "pictures": walked["pictures"],
        "shapes": walked["shapes"],
    }


def main():
    parser = argparse.ArgumentParser(description="Inspect PPTX shape/font/color properties")
    parser.add_argument("--pptx", required=True)
    parser.add_argument("--slide", type=int, default=0)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    path = Path(args.pptx)
    if not path.exists():
        print(f"ERROR: PPTX not found: {path}")
        sys.exit(1)
    data = inspect_pptx(path, args.slide)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, indent=2), encoding="utf-8")
    print(
        f"Inspected {len(data['shapes'])} objects "
        f"({data['groups']} groups, {data['pictures']} pictures) → {out}"
    )


if __name__ == "__main__":
    main()
