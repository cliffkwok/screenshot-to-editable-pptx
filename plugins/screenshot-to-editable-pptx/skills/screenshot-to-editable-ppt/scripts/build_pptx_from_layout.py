#!/usr/bin/env python3
"""Build a PPTX from a universal layout JSON (source-pixel coordinates).

Schema (minimal):
{
  "source_width": 1024,
  "source_height": 635,
  "slide_h_in": 7.5,
  "background": "#FFFFFF",
  "elements": [
    {"type": "text", "name": "title", "box": [x,y,w,h], "text": "...",
     "font_pt": 18, "bold": true, "color": "#222", "align": "left"},
    {"type": "shape", "name": "card", "shape": "rounded_rect",
     "box": [x,y,w,h], "fill": "#fff", "line": "#ccc", "line_pt": 0.75, "rad": 0.1},
    {"type": "line", "name": "div", "points": [x1,y1,x2,y2],
     "color": "#ddd", "line_pt": 0.75},
    {"type": "image", "name": "icon1", "box": [x,y,w,h],
     "path": "assets/icon1.png", "png_fallback": "assets/icon1.png", "svg": "assets/icon1.svg"}
  ]
}

z-order = list order (background first). Prefer PNG for visible image when
svg is present (png-fallback policy).

Usage:
  python3 scripts/build_pptx_from_layout.py \\
      --layout work/deck/layout.json --assets-root work/deck --out work/deck/out.pptx
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from helpers import (  # noqa: E402
    Inches, Pt, add_line, add_shape, add_shape_with_text, text_in_box,
    new_presentation, set_background,
)

SHAPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "ellipse": MSO_SHAPE.OVAL,
    "diamond": MSO_SHAPE.DIAMOND,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "trapezoid": MSO_SHAPE.TRAPEZOID,
}

DASH_MAP = {
    "dash": MSO_LINE_DASH_STYLE.DASH,
    "dash_dot": MSO_LINE_DASH_STYLE.DASH_DOT,
    "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
}


def resolve_asset(path: str | None, assets_root: Path) -> Path | None:
    if not path:
        return None
    p = Path(path)
    if not p.is_absolute():
        p = assets_root / p
    return p if p.exists() else None


def build_from_layout(layout: dict, assets_root: Path, out: Path) -> None:
    img_w = int(layout["source_width"])
    img_h = int(layout["source_height"])
    slide_h = float(layout.get("slide_h_in") or 7.5)
    slide_w = float(layout.get("slide_w_in") or round(slide_h * img_w / img_h, 3))
    scale = slide_h / img_h
    bg = layout.get("background") or "#FFFFFF"

    def px(x, y):
        return Inches(x * scale), Inches(y * scale)

    def inch_box(x, y, w, h):
        return (
            Inches(x * scale), Inches(y * scale),
            Inches(w * scale), Inches(h * scale),
        )

    prs, slide = new_presentation(slide_w, slide_h)
    # Solid page color = slide.background only. Never add a full-slide rect.
    set_background(slide, bg)

    elements = layout.get("elements")
    if not elements and layout.get("slides"):
        # single-slide: use first slide only (multi-page: call once per page)
        elements = layout["slides"][0].get("elements") or []
    if not elements:
        raise ValueError("layout has no elements")

    for el in elements:
        typ = (el.get("type") or "shape").lower()
        name = el.get("name") or el.get("id") or typ
        box = el.get("box") or el.get("box_px")
        # Skip redundant full-slide backing shapes (same as layout.background)
        if name in ("slide_bg", "bg", "background") and box:
            x, y, w, h = box[:4]
            if (x <= 1 and y <= 1
                    and abs(w - img_w) <= 2 and abs(h - img_h) <= 2):
                print(f"  skip full-slide backing shape {name!r} (use background hex)")
                continue

        if typ in ("image", "picture"):
            # png-fallback: visible layer prefers PNG when SVG also listed
            png = resolve_asset(el.get("png_fallback") or el.get("path"), assets_root)
            if png is None:
                png = resolve_asset(el.get("path"), assets_root)
            if png is None:
                print(f"  WARN skip image {name}: missing file")
                continue
            if not box or len(box) != 4:
                print(f"  WARN skip image {name}: no box")
                continue
            x, y, w, h = box
            left, top, width, height = inch_box(x, y, w, h)
            pic = slide.shapes.add_picture(str(png), left, top, width=width, height=height)
            pic.name = name
            continue

        if typ == "line":
            pts = el.get("points")
            if not pts or len(pts) < 4:
                print(f"  WARN skip line {name}")
                continue
            x1, y1, x2, y2 = pts[:4]
            col = el.get("color") or el.get("line") or "#000000"
            lw = float(el.get("line_pt") or el.get("width_pt") or 0.75)
            ln = add_line(
                slide, *px(x1, y1), *px(x2, y2),
                color=col, width=Pt(lw), kind=el.get("kind") or "straight",
            )
            dash = el.get("dash")
            if dash and dash in DASH_MAP:
                ln.line.dash_style = DASH_MAP[dash]
            ln.name = name
            continue

        if typ in ("text", "textbox"):
            if not box or len(box) != 4:
                print(f"  WARN skip text {name}")
                continue
            text = el.get("text") or ""
            size = float(el.get("font_pt") or el.get("size") or 12)
            sh = text_in_box(
                slide, box, text, size, img_w, img_h, slide_w, slide_h,
                bold=bool(el.get("bold")),
                color=el.get("color") or "#1A1A1A",
                align=el.get("align") or "left",
                font=el.get("font") or "Arial",
                pad=0,
                wrap=bool(el.get("wrap")),
            )
            sh.name = name
            continue

        # shape
        if not box or len(box) != 4:
            print(f"  WARN skip shape {name}")
            continue
        shape_key = (el.get("shape") or "rect").lower()
        mso = SHAPE_MAP.get(shape_key, MSO_SHAPE.RECTANGLE)
        fill = el.get("fill")
        if fill in ("none", "transparent", ""):
            fill = None
        line = el.get("line")
        if line in ("none", "null", ""):
            line = None
        lw = el.get("line_pt")
        lw = Pt(float(lw)) if lw is not None else (Pt(0.75) if line else None)
        rad = el.get("rad") or el.get("radius")
        text = el.get("text")
        if text:
            sh = add_shape_with_text(
                slide, mso, *inch_box(*box), text,
                float(el.get("font_pt") or 12),
                fill=fill, line=line, lw=lw, rad=rad,
                bold=bool(el.get("bold")),
                color=el.get("color") or "#1A1A1A",
                font=el.get("font") or "Arial",
                align=el.get("align") or "center",
                anchor="middle", margin_pt=2, fit_to_shape=False,
                wrap=bool(el.get("wrap")),
            )
        else:
            sh = add_shape(
                slide, mso, *inch_box(*box),
                fill=fill, line=line, lw=lw, rad=rad,
            )
        sh.name = name

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out}  ({len(elements)} elements, {slide_w}×{slide_h} in)")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--layout", required=True)
    ap.add_argument("--assets-root", default=".")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    layout = json.loads(Path(args.layout).read_text())
    build_from_layout(layout, Path(args.assets_root), Path(args.out))


if __name__ == "__main__":
    main()
