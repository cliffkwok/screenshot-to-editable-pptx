"""python-pptx helpers for screenshot reconstruction.

Import from a builder script after putting this folder on sys.path:

    from helpers import (
        new_presentation, add_shape, add_text, add_shadow,
        disable_shadow, group_shapes, rgb, inch_pct
    )
"""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mapping import (  # noqa: E402
    Image,
    expected_glyph_h_px,
    font_pt_fit_shape,
    font_pt_from_glyph,
    ink_inside_shape,
    measure_text_in_shape,
    pad_ink_box_px,
    pct_to_px,
    px_box_to_inches,
    screenshot_mapping,
)

ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def new_presentation(width_in=13.33, height_in=7.5):
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def set_background(slide, hex_color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)


def disable_shadow(shape):
    """Kill python-pptx default shadow. Call on every shape."""
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    return shape


def add_shape(slide, typ, l, t, w, h, fill=None, line=None, lw=None, rad=None):
    sh = slide.shapes.add_shape(typ, l, t, w, h)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill) if isinstance(fill, str) else fill
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line) if isinstance(line, str) else line
        sh.line.width = lw or Pt(1)
    else:
        sh.line.fill.background()
    if rad is not None:
        sh.adjustments[0] = rad
    disable_shadow(sh)
    return sh


def set_shape_text(
    shape, text, size_pt, bold=False, color="#FFFFFF", font="Arial",
    align="center", anchor="middle", margin_pt=3, autofit=False, wrap=False,
    fit_to_shape=True, circle=False, pad_frac=None,
):
    """Put text *inside* a shape (PowerPoint: select shape → Edit Text).

    Prefer this over a separate text box when the screenshot shows a label
    on a filled bar/pill/circle. Point size is capped so the line box stays
    inside the shape with the same gap as the photo (`pad_frac` measured from
    ink vs fill). `autofit` is off by default so PPT does not fight the size.
    """
    pt = float(size_pt)
    if fit_to_shape:
        pt = font_pt_fit_shape(
            text,
            shape.width.inches,
            shape.height.inches,
            glyph_pt=pt,
            margin_pt=margin_pt,
            circle=circle,
            pad_frac=pad_frac,
        )
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if autofit else MSO_AUTO_SIZE.NONE
    if anchor == "middle":
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif anchor == "bottom":
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    else:
        tf.vertical_anchor = MSO_ANCHOR.TOP
    m = Pt(margin_pt)
    tf.margin_left = m
    tf.margin_right = m
    tf.margin_top = m
    tf.margin_bottom = m
    p = tf.paragraphs[0]
    p.alignment = ALIGN.get(align, PP_ALIGN.CENTER)
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(pt)
    run.font.bold = bold
    run.font.color.rgb = rgb(color) if isinstance(color, str) else color
    run.font.name = font
    shape._fit_font_pt = pt
    return shape


def add_shape_with_text(
    slide, typ, l, t, w, h, text, size_pt,
    fill=None, line=None, lw=None, rad=None,
    bold=False, color="#FFFFFF", font="Arial",
    align="center", anchor="middle", margin_pt=3, autofit=False,
    fit_to_shape=True, circle=False, pad_frac=None, wrap=False,
):
    """Create a shape and put `text` in its text frame (not a floating text box)."""
    sh = add_shape(slide, typ, l, t, w, h, fill=fill, line=line, lw=lw, rad=rad)
    set_shape_text(
        sh, text, size_pt, bold=bold, color=color, font=font,
        align=align, anchor=anchor, margin_pt=margin_pt, autofit=autofit,
        fit_to_shape=fit_to_shape, circle=circle, pad_frac=pad_frac, wrap=wrap,
    )
    return sh


def add_line(slide, x1, y1, x2, y2, color="#000000", width=Pt(1.5), dash=None,
             begin_arrow=None, end_arrow=None, arrow_size="sm"):
    """Real PPT connector line (not a thin rectangle or arrow AutoShape)."""
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    sh.line.color.rgb = rgb(color) if isinstance(color, str) else color
    sh.line.width = width
    if dash is not None:
        sh.line.dash_style = dash
    set_line_arrows(sh, begin=begin_arrow, end=end_arrow, size=arrow_size)
    disable_shadow(sh)
    return sh


# OOXML arrowhead types shown in Format Shape → Line → Begin/End Arrow
_ARROW_TYPE = {
    None: None,
    "none": None,
    "triangle": "triangle",
    "arrow": "triangle",
    "stealth": "stealth",
    "diamond": "diamond",
    "oval": "oval",
    "open": "arrow",
}
_ARROW_SIZE = {"sm": "sm", "small": "sm", "med": "med", "medium": "med",
               "lg": "lg", "large": "lg"}


def set_line_arrows(shape, begin=None, end=None, size="sm"):
    """Set PowerPoint line arrowheads (Format Shape → Line → Begin/End Arrow).

    Never fake an arrow with MSO_SHAPE.RIGHT_ARROW. `begin` is headEnd
    (start of the connector); `end` is tailEnd.
    """
    ln = shape.line._get_or_add_ln()
    for tag in ("a:headEnd", "a:tailEnd"):
        el = ln.find(qn(tag))
        if el is not None:
            ln.remove(el)
    sz = _ARROW_SIZE.get(size, "sm")
    b = _ARROW_TYPE.get(begin, begin)
    e = _ARROW_TYPE.get(end, end)
    if b:
        etree.SubElement(ln, qn("a:headEnd"), {"type": str(b), "w": sz, "len": sz})
    if e:
        etree.SubElement(ln, qn("a:tailEnd"), {"type": str(e), "w": sz, "len": sz})
    return shape


def stroke_px_to_pt(stroke_px, img_h, slide_h_in=7.5):
    """Screenshot stroke thickness → line width in points."""
    if not stroke_px or img_h <= 0:
        return 1.0
    return max(0.5, round(float(stroke_px) / float(img_h) * float(slide_h_in) * 72.0, 2))


def measure_stroke_px(im, x, y, axis="v", dark=200, span=18):
    """Perpendicular run-length of dark ink at (x,y). axis='v' for a horizontal line."""
    if Image is None:
        raise RuntimeError("Pillow is required")
    rgb = im.convert("RGB") if im.mode != "RGB" else im
    px = rgb.load()
    w, h = rgb.size

    def darkish(xx, yy):
        if not (0 <= xx < w and 0 <= yy < h):
            return False
        r, g, b = px[xx, yy]
        return (r + g + b) / 3 < dark

    best = 0
    if axis == "v":
        for x0 in range(max(0, x - 6), min(w, x + 7)):
            run = 0
            for yy in range(max(0, y - span), min(h, y + span + 1)):
                if darkish(x0, yy):
                    run += 1
                    best = max(best, run)
                else:
                    run = 0
    else:
        for y0 in range(max(0, y - 6), min(h, y + 7)):
            run = 0
            for xx in range(max(0, x - span), min(w, x + span + 1)):
                if darkish(xx, y0):
                    run += 1
                    best = max(best, run)
                else:
                    run = 0
    return best


# Connection sites on a shape bounding box (python-pptx convention):
# 0 = top, 1 = left, 2 = bottom, 3 = right
CXN_TOP, CXN_LEFT, CXN_BOTTOM, CXN_RIGHT = 0, 1, 2, 3
_CONNECTOR_KIND = {
    "straight": MSO_CONNECTOR.STRAIGHT,
    "elbow": MSO_CONNECTOR.ELBOW,
    "curve": MSO_CONNECTOR.CURVE,
}


def connect_shapes(
    slide, start_shape, start_idx, end_shape, end_idx,
    color="#000000", width=Pt(1.5), kind="straight", dash=None,
    begin_arrow=None, end_arrow=None, arrow_size="sm",
):
    """Attach a connector to two shapes so it stays connected in PowerPoint.

    Circles use the same 4 bbox sites; those land on N/E/S/W of the oval.
    Create the two shapes first, then this connector. Arrowheads are line
    properties (Format Shape → Line), not arrow AutoShapes.

    `kind`: detect from the screenshot — "straight", "elbow" (right angles),
    or "curve" (smooth arcs; PowerPoint Curved Connector / Curve look).
    Do not default to elbow when the photo shows smooth bends.
    """
    def _pt(shape, idx):
        x, y, w, h = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
        return {
            0: (x + w // 2, y),
            1: (x, y + h // 2),
            2: (x + w // 2, y + h),
            3: (x + w, y + h // 2),
        }[idx]

    x1, y1 = _pt(start_shape, start_idx)
    x2, y2 = _pt(end_shape, end_idx)
    sh = slide.shapes.add_connector(_CONNECTOR_KIND.get(kind, MSO_CONNECTOR.STRAIGHT), x1, y1, x2, y2)
    sh.begin_connect(start_shape, start_idx)
    sh.end_connect(end_shape, end_idx)
    sh.line.color.rgb = rgb(color) if isinstance(color, str) else color
    sh.line.width = width
    if dash is not None:
        sh.line.dash_style = dash
    set_line_arrows(sh, begin=begin_arrow, end=end_arrow, size=arrow_size)
    disable_shadow(sh)
    return sh


def connect_lr(slide, a, b, color="#000000", width=Pt(1.5), kind="straight", dash=None,
               begin_arrow=None, end_arrow=None, arrow_size="sm"):
    """Connect a's right site to b's left site."""
    return connect_shapes(
        slide, a, CXN_RIGHT, b, CXN_LEFT, color, width, kind, dash,
        begin_arrow=begin_arrow, end_arrow=end_arrow, arrow_size=arrow_size,
    )


def inset_box(box, pad=0, pad_l=None, pad_t=None, pad_r=None, pad_b=None):
    """Shrink a pixel box by padding. Uniform `pad` or per-side overrides."""
    x, y, w, h = [float(v) for v in box]
    l = pad if pad_l is None else pad_l
    t = pad if pad_t is None else pad_t
    r = pad if pad_r is None else pad_r
    b = pad if pad_b is None else pad_b
    return (x + l, y + t, max(1.0, w - l - r), max(1.0, h - t - b))


def center_in(parent, child_w, child_h):
    """Pixel box for a child of size (child_w, child_h) centered in parent."""
    px, py, pw, ph = [float(v) for v in parent]
    return (px + (pw - child_w) / 2.0, py + (ph - child_h) / 2.0, child_w, child_h)


def band(parent, y0_frac, y1_frac):
    """Horizontal band inside parent from y0_frac→y1_frac of parent height."""
    x, y, w, h = [float(v) for v in parent]
    t = y + h * y0_frac
    b = y + h * y1_frac
    return (x, t, w, max(1.0, b - t))


def label_above(cx, cy, node_r, text_w, text_h, gap=6):
    """Label box centered above a circular node."""
    return (cx - text_w / 2.0, cy - node_r - gap - text_h, text_w, text_h)


def label_below(cx, cy, node_r, text_w, text_h, gap=6):
    """Label box centered below a circular node."""
    return (cx - text_w / 2.0, cy + node_r + gap, text_w, text_h)


def add_line_text(slide, l, t, w, h, text, size_pt, bold=False, color="#1A1A1A",
                  align="left", font="Arial", anchor="middle"):
    """One visual line. wrap=False so PPT never wraps a measured glyph box."""
    return add_text(
        slide, l, t, w, h, text, size_pt, bold=bold, color=color,
        align=align, font=font, auto_fit=False, anchor=anchor, wrap=False,
    )


def text_in_box(slide, box_px, text, size_pt, img_w, img_h, slide_w=13.33, slide_h=7.5,
                fit="height", bold=False, color="#1A1A1A", align="center", font="Arial",
                pad=2, wrap=False, anchor="middle"):
    """Place text in a container box (parent inset).

    `wrap=False` (default): one visual line — measured labels, titles.
    `wrap=True`: PowerPoint placeholder flow — paragraph wraps inside the box
    width (captions that break to the next line in the screenshot).
    """
    x, y, w, h = inset_box(box_px, pad=pad)
    l, t, ww, hh = px_box_to_inches(x, y, w, h, img_w, img_h, slide_w, slide_h, fit)
    if wrap:
        return add_text(
            slide, l, t, ww, hh, text, size_pt,
            bold=bold, color=color, align=align, font=font,
            auto_fit=False, anchor=anchor, wrap=True,
        )
    return add_line_text(
        slide, l, t, ww, hh, text, size_pt,
        bold=bold, color=color, align=align, font=font, anchor=anchor,
    )


def caption_band_box(card_box, diagram_box, pad_top_px, pad_bot_px, side_pad_px=10):
    """Placeholder box for a caption under a diagram, from measured gaps.

    Measure on the screenshot:
      pad_top = distance from diagram bottom edge → caption ink top
      pad_bot = distance from caption ink bottom → next element / card bottom
    Width ≈ card width − 2×side_pad (or measured ink width + side pad).
    Alignment is usually center for card footers; left for callouts beside nodes.
    """
    cx, cy, cw, ch = [float(v) for v in card_box]
    _dx, _dy, _dw, dh = [float(v) for v in diagram_box]
    diag_bottom = _dy + dh
    y = diag_bottom + float(pad_top_px)
    # Prefer explicit bottom pad from card bottom when provided
    h = (cy + ch - float(pad_bot_px)) - y
    if h < 12:
        h = 12.0
    x = cx + float(side_pad_px)
    w = max(20.0, cw - 2 * float(side_pad_px))
    return (x, y, w, h)


def abs_in_parent(parent_box, local_xywh):
    """Convert parent-local (x,y,w,h) → absolute screenshot/slide px box."""
    px, py, pw, ph = [float(v) for v in parent_box]
    lx, ly, lw, lh = [float(v) for v in local_xywh]
    return (px + lx, py + ly, lw, lh)


def child_fits(parent_box, child_box, slack_px=2):
    """True if child is inside parent (optional slack for AA / intentional overlap)."""
    px, py, pw, ph = [float(v) for v in parent_box]
    cx, cy, cw, ch = [float(v) for v in child_box]
    return (
        cx >= px - slack_px
        and cy >= py - slack_px
        and cx + cw <= px + pw + slack_px
        and cy + ch <= py + ph + slack_px
    )


def assert_children_inside(parent_box, children, slack_px=2):
    """Raise if any named child overflows the nested parent.

    `children` is [{name, box_px}, ...]. Use after measuring layout and
    BEFORE grouping — overflow inside a group is what PowerPoint shows as
    piled / overlapping members when the group is selected.
    """
    bad = []
    for ch in children:
        box = ch.get("box_px") or ch.get("box")
        name = ch.get("name") or ch.get("id") or "?"
        if box is None:
            continue
        if not child_fits(parent_box, box, slack_px=slack_px):
            bad.append(f"{name} box={list(map(int, box))} ⊄ parent={list(map(int, parent_box))}")
    if bad:
        raise ValueError(
            "nested children overflow parent (fix layout / nest pads):\n  - "
            + "\n  - ".join(bad)
        )


def point_on_edge(box, side, t=0.5):
    """Pixel point on a box edge. `t` in [0,1] along the edge (0=start).

    sides: left / right / top / bottom. Used so connectors terminate ON the
    stroke, not at a shared center that piles every spoke on one pixel.
    """
    x, y, w, h = [float(v) for v in box]
    t = max(0.0, min(1.0, float(t)))
    if side in ("left", "l"):
        return (x, y + h * t)
    if side in ("right", "r"):
        return (x + w, y + h * t)
    if side in ("top", "t"):
        return (x + w * t, y)
    if side in ("bottom", "b"):
        return (x + w * t, y + h)
    raise ValueError(f"unknown side {side!r}")


def edge_attach_ts(n, margin=0.12):
    """Evenly spaced `t` values along an edge for n fan-in/out spokes.

    Avoids the classic defect where every connector ends at the mid-point
    and the strokes look like one thick overlapping bundle.
    """
    n = int(n)
    if n <= 0:
        return []
    if n == 1:
        return [0.5]
    lo, hi = float(margin), 1.0 - float(margin)
    if n == 2:
        return [lo, hi]
    return [lo + (hi - lo) * i / (n - 1) for i in range(n)]


def uniform_stack(y0, h, gap, n):
    """n boxes of height h with constant gap — sibling rows from photo measure."""
    return [float(y0) + i * (float(h) + float(gap)) for i in range(int(n))]


def assert_min_gap(a_box, b_box, axis, min_px, name=""):
    """Raise if two axis-aligned boxes are closer than min_px (or overlap)."""
    ax, ay, aw, ah = [float(v) for v in a_box]
    bx, by, bw, bh = [float(v) for v in b_box]
    if axis == "x":
        if ax + aw <= bx:
            gap = bx - (ax + aw)
        elif bx + bw <= ax:
            gap = ax - (bx + bw)
        else:
            gap = -1.0
    elif axis == "y":
        if ay + ah <= by:
            gap = by - (ay + ah)
        elif by + bh <= ay:
            gap = ay - (by + bh)
        else:
            gap = -1.0
    else:
        raise ValueError("axis must be 'x' or 'y'")
    if gap < float(min_px):
        tag = f" ({name})" if name else ""
        raise ValueError(
            f"gap too small{tag}: {gap:.1f}px < {min_px}px "
            f"a={list(map(int, a_box))} b={list(map(int, b_box))}"
        )


def add_flow_text(slide, l, t, w, h, text, size_pt, bold=False, color="#1A1A1A",
                  align="center", font="Arial", anchor="middle"):
    """One text-box placeholder: word wrap on, text flows to the next line.

    Use when the screenshot shows a wrapped paragraph (caption under a card),
    not when each line is a separately measured label.
    """
    return add_text(
        slide, l, t, w, h, text, size_pt,
        bold=bold, color=color, align=align, font=font,
        auto_fit=False, anchor=anchor, wrap=True,
    )


def add_text_from_hint(
    slide, hint, text=None, color="#1A1A1A", font="Arial", bold=False,
    script=None, img_w=None, img_h=None, slide_w=13.33, slide_h=7.5,
    fit="height", align="left",
):
    """Place a single-line text box from a text_hints.py record."""
    script = script or hint.get("script") or "latin"
    key = "size_group_font_pt_cjk" if script == "cjk" else "size_group_font_pt_latin"
    size = hint[key]
    x, y, w, h = hint["box_px"]
    x, y, w, h = pad_ink_box_px(x, y, w, h, img_w=img_w, img_h=img_h)
    l, t, ww, hh = px_box_to_inches(x, y, w, h, img_w, img_h, slide_w, slide_h, fit)
    return add_line_text(
        slide, l, t, ww, hh, text if text is not None else hint.get("text", ""),
        size, bold=bold, color=color, align=align, font=font,
    )


def add_text(slide, l, t, w, h, text, size_pt, bold=False, color="#1A1A1A",
             align="center", font="Arial", auto_fit=False, anchor="middle",
             wrap=True):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    if auto_fit:
        tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    if anchor == "middle":
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif anchor == "bottom":
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    else:
        tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = ALIGN.get(align, PP_ALIGN.CENTER)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = font
    disable_shadow(tb)
    return tb


def add_shadow(shape, blur=25000, dist=8000, alpha=8000, base="999999"):
    """Apply outer shadow. Only when Vision marked the element as having one."""
    spPr = shape._element.find(qn("a:spPr"))
    if spPr is None:
        spPr = shape._element
    el = spPr.find(qn("a:effectLst"))
    if el is None:
        el = etree.SubElement(spPr, qn("a:effectLst"))
    existing = el.find(qn("a:outerShdw"))
    if existing is not None:
        el.remove(existing)
    os = etree.SubElement(el, qn("a:outerShdw"), {
        "blurRad": str(blur),
        "dist": str(dist),
        "dir": "5400000",
        "rotWithShape": "0",
    })
    sc = etree.SubElement(os, qn("a:srgbClr"), {"val": base})
    etree.SubElement(sc, qn("a:alpha"), {"val": str(alpha)})
    return shape


def inch_pct(container_origin, container_size, pct):
    """container origin/size as Inches; pct is 0-100. Returns Inches."""
    return container_origin + container_size * pct / 100.0


def cjk_em_in(size_pt: float) -> float:
    """Full-width CJK advance at this point size, in inches (1em)."""
    return float(size_pt) / 72.0


def cjk_centered_substring_box(
    text: str,
    substring: str,
    size_pt: float,
    slide_w_in: float,
    box_top_in: float,
    box_h_in: float,
    pad_x_in: float = 0.03,
    pad_y_in: float = 0.04,
):
    """Slide-absolute Inches (l, t, w, h) for a marker behind a CJK substring.

    The title is assumed centered on the slide. Never measure the marker as a
    second screenshot blob — proofing squiggles and compression inflate it.
    """
    if substring not in text:
        raise ValueError(f"{substring!r} not in {text!r}")
    em = cjk_em_in(size_pt)
    total = len(text) * em
    left = (slide_w_in - total) / 2.0
    i = text.index(substring)
    l = left + i * em - pad_x_in
    w = len(substring) * em + 2 * pad_x_in
    h = min(max(box_h_in - 0.08, em * 0.7), em * 0.92)
    t = box_top_in + (box_h_in - h) / 2.0
    return Inches(l), Inches(t), Inches(w), Inches(h)


def _xfrm_of(el):
    for child in el:
        if child.tag in (qn("p:spPr"), qn("p:grpSpPr")):
            xf = child.find(qn("a:xfrm"))
            if xf is not None:
                return xf
    return el.find(".//" + qn("a:xfrm"))


def _next_cNvPr_id(spTree) -> int:
    ids = [0]
    for n in spTree.iter(qn("p:cNvPr")):
        try:
            ids.append(int(n.get("id", 0)))
        except (TypeError, ValueError):
            pass
    return max(ids) + 1


def add_pill_icon(
    slide, l, t, w, h, stroke,
    fill="#FFFFFF", lw=None, eye_frac=0.26, eye_x=(0.32, 0.68), name="pill",
):
    """Capsule (rounded rect) + two eye dots → one PowerPoint group.

    Composite icons that read as a single glyph (pill shell + inner dots)
    must be grouped. Select once in PowerPoint = one bounding box.
    Connectors stay *outside* this group and attach to the group shape.

    `lw` is the capsule outline weight — usually heavier than diagram
    connectors (measure both on the screenshot; do not reuse one pt).
    """
    if lw is None:
        lw = Pt(2.5)
    body = add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h,
        fill=fill, line=stroke, lw=lw, rad=0.5,
    )
    eye = h * eye_frac
    ey = t + (h - eye) / 2
    dots = []
    for fx in eye_x:
        dots.append(add_shape(
            slide, MSO_SHAPE.OVAL,
            l + w * fx - eye / 2, ey, eye, eye,
            fill=stroke, line=None,
        ))
    return group_shapes(slide, [body, *dots], name=name)


def group_shapes(slide, shapes, name="Group", nest_groups=False):
    """Group shapes via OOXML so the result scales in PowerPoint.

    PowerPoint world coords are:
        world = grp.off + (child.off - chOff) * (ext / chExt)

    Children MUST be local (origin 0) and chOff MUST be (0, 0). Setting
    chOff to the group's slide origin AND subtracting that origin from
    children (or assigning shape.left/top after the node is inside grpSp)
    double-offsets the group. PowerPoint then parks it at the slide origin
    even though python-pptx .left still looks correct.

    By default, existing groups (e.g. a pill icon) stay on the slide as
    their own groups. Nesting them inside a card group makes PowerPoint
    click-select the inner rounded rect instead of the icon group.
    Pass nest_groups=True only when you truly need a group-of-groups.
    """
    if not shapes:
        return None
    members = []
    for s in shapes:
        if s is None:
            continue
        try:
            if (not nest_groups) and getattr(s, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                continue
        except Exception:
            pass
        members.append(s)
    if not members:
        return None
    spTree = slide.shapes._spTree
    records = [(s, int(s.left), int(s.top), int(s.width), int(s.height)) for s in members]
    min_x = min(r[1] for r in records)
    min_y = min(r[2] for r in records)
    max_x = max(r[1] + r[3] for r in records)
    max_y = max(r[2] + r[4] for r in records)
    width = max_x - min_x
    height = max_y - min_y

    grpSp = etree.SubElement(spTree, qn("p:grpSp"))
    nvGrpSpPr = etree.SubElement(grpSp, qn("p:nvGrpSpPr"))
    etree.SubElement(nvGrpSpPr, qn("p:cNvPr"), {
        "id": str(_next_cNvPr_id(spTree)),
        "name": name,
    })
    etree.SubElement(nvGrpSpPr, qn("p:cNvGrpSpPr"))
    etree.SubElement(nvGrpSpPr, qn("p:nvPr"))

    grpSpPr = etree.SubElement(grpSp, qn("p:grpSpPr"))
    xfrm = etree.SubElement(grpSpPr, qn("a:xfrm"))
    etree.SubElement(xfrm, qn("a:off"), {"x": str(min_x), "y": str(min_y)})
    etree.SubElement(xfrm, qn("a:ext"), {"cx": str(width), "cy": str(height)})
    etree.SubElement(xfrm, qn("a:chOff"), {"x": "0", "y": "0"})
    etree.SubElement(xfrm, qn("a:chExt"), {"cx": str(width), "cy": str(height)})

    for s, left, top, _w, _h in records:
        sp = s._element
        if sp is grpSp:
            continue
        spTree.remove(sp)
        grpSp.append(sp)
        # XML only. Never assign s.left / s.top after the node is inside grpSp.
        xf = _xfrm_of(sp)
        if xf is None:
            raise RuntimeError(f"group {name}: {s.name} has no a:xfrm")
        off = xf.find(qn("a:off"))
        if off is None:
            raise RuntimeError(f"group {name}: {s.name} has no a:off")
        off.set("x", str(left - min_x))
        off.set("y", str(top - min_y))
    # Return the python-pptx GroupShape wrapper (has .left/.top for connectors).
    for shape in slide.shapes:
        if shape._element is grpSp:
            return shape
    return grpSp


# Re-export common enums so builders can `from helpers import MSO_SHAPE, Inches, Pt`
__all__ = [
    "MSO_SHAPE", "Inches", "Pt", "Emu", "PP_ALIGN", "MSO_ANCHOR", "MSO_AUTO_SIZE",
    "rgb", "new_presentation", "set_background", "disable_shadow",
    "add_shape", "add_shape_with_text", "set_shape_text", "add_text", "add_flow_text",
    "add_line", "add_line_text", "add_text_from_hint",
    "connect_shapes", "connect_lr", "set_line_arrows", "stroke_px_to_pt", "measure_stroke_px",
    "CXN_TOP", "CXN_LEFT", "CXN_BOTTOM", "CXN_RIGHT",
    "add_shadow", "add_pill_icon",
    "inch_pct", "group_shapes", "cjk_em_in", "cjk_centered_substring_box",
    "font_pt_from_glyph", "font_pt_fit_shape", "measure_text_in_shape",
    "ink_inside_shape", "expected_glyph_h_px", "pad_ink_box_px",
    "Image",
    "px_box_to_inches", "pct_to_px", "screenshot_mapping",
    "inset_box", "center_in", "band", "label_above", "label_below", "text_in_box",
    "caption_band_box",
    "abs_in_parent", "child_fits", "assert_children_inside",
    "point_on_edge", "edge_attach_ts", "uniform_stack", "assert_min_gap",
]
